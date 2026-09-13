# -*- coding: utf-8 -*-
"""Apply the 2026-09 revision to the manuscript: text, tables and the
regenerated route figure.  Input: docs/manuscript_base.docx -> out/."""
import os as _os
ROOT = _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
FIGS = _os.environ.get("STUDIO_FIGS", _os.path.join(ROOT, "studies", "out", "figs"))
OUT = _os.path.join(ROOT, "studies", "out"); _os.makedirs(OUT, exist_ok=True)
from docx import Document
doc = Document(_os.path.join(ROOT, "docs", "manuscript_base.docx"))
P = doc.paragraphs
def find(sub):
    for i, p in enumerate(P):
        if sub in p.text: return i
def rep(sub, new):
    i = find(sub)
    if i is None: print("MISS:", sub[:50]); return
    for r in P[i].runs:
        if sub in r.text: r.text = r.text.replace(sub, new); return
    t = P[i].text.replace(sub, new)
    for r in P[i].runs[1:]: r.text = ""
    P[i].runs[0].text = t
rep("live-load deflection limit of span/360 on a cracked effective inertia",
    "live-load deflection limit of span/360 on the ACI 318-19 effective inertia (Bischoff form) of the cracked transformed section, with the d/4 stirrup-spacing rule when Vs exceeds 0.33\u221af\u2032c\u00b7b\u00b7d")
rep("with lateral-torsional buckling per the governing chapter at Cb = 1.",
    "with lateral-torsional buckling per the governing chapter and the modification factor Cb an explicit bounded input (default 1.0; 1.14 for a simply supported uniform load).")
rep("the built-up beam module treats unstiffened doubly symmetric sections at Cb = 1;",
    "the built-up beam module treats unstiffened doubly symmetric sections;")
rep("The beam optimum (728.3 kg CO", "The beam optimum (727.8 kg CO")
rep("cuts the optimal footprint by 56% for the beam (728 \u2192 323 kg) while its optimal steel mass rises 57% (140 \u2192 220 kg) in a shallower section",
    "cuts the optimal footprint by 56% for the beam (728 \u2192 321 kg) while its optimal steel mass rises 88% (136 \u2192 255 kg) in a shallower section")
i = find("assign extensions")
if i is not None:
    P[i].runs[-1].text += " A six-module capacity-building pathway\u2014learning outcomes, seeded exercises and expected outcomes drawn from the studies of Section 4\u2014accompanies the code for course and workshop use."
i = find("26-test verification suite")
if i is not None:
    P[i].runs[-1].text = P[i].runs[-1].text.replace("suite that pins every mechanical anchor used in this paper.",
        "suite that pins every mechanical anchor used in this paper, and a six-module capacity-building curriculum.")
def setcell(cell, text):
    for k, p in enumerate(cell.paragraphs):
        for j, r in enumerate(p.runs): r.text = text if (k == 0 and j == 0) else ""
T = doc.tables
for row in T[0].rows:
    if row.cells[0].text.strip() == "RC beam": setcell(row.cells[-1], "727.8")
for row in T[2].rows:
    if row.cells[0].text.strip() == "RC beam":
        for c, t in zip(row.cells, ["RC beam", "728.3 kg CO\u2082", "727.8", "+0.1%", "flexure (100%)"]): setcell(c, t)
beamvals = {"BEDEC placed rebar": "728 kg (136 kg)", "BF\u2013BOF primary": "609 kg (150 kg)",
            "World average": "587 kg (154 kg)", "EAF scrap, grid": "386 kg (178 kg)",
            "EAF, low-carbon": "321 kg (255 kg)"}
for row in T[3].rows:
    k = row.cells[0].text.strip()
    if k in beamvals: setcell(row.cells[2], beamvals[k])
# replace the Fig. 6 image with the regenerated route figure
ic = find("Fig. 6. Route study")
img_p = P[ic-1]
blips = img_p._p.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
rid = blips[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
part = doc.part.related_parts[rid]
part._blob = open(_os.path.join(FIGS, "fig_routes.png"), "rb").read()
print("Fig. 6 image replaced via", rid)
doc.save(_os.path.join(OUT, "Structural_Optimization_Manuscript.docx"))
print("manuscript written")

# =========================== revision 2: carbon everywhere, education, GUI, flow
from docx.shared import Emu
from PIL import Image as _Image
def swap_figure(caption_key, png, width_px):
    ic = find(caption_key); img_p = P[ic-1]
    blip = img_p._p.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')[0]
    rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
    doc.part.related_parts[rid]._blob = open(png, "rb").read()
    w, h = _Image.open(png).size
    for sh in doc.inline_shapes:
        b2 = sh._inline.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
        if b2 and b2[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed') == rid:
            sh.width = Emu(int(width_px*9525)); sh.height = Emu(int(width_px*h/w*9525))
    print("swapped", caption_key[:8], "->", png.split("/")[-1])
swap_figure("Fig. 1. Framework architecture", _os.path.join(FIGS, "fig_flow.png"), 470)
swap_figure("Fig. 2. Interactive workspace", _os.path.join(FIGS, "fig_gui.png"), 560)
rep("Fig. 1. Framework architecture: formulations, analysis, carbon accounting, optimizer and interactive layer.",
    "Fig. 1. System flowchart: from the problem library and user inputs through the genetic-algorithm loop (evaluation, fitness mapping, reproduction) to post-processing with carbon accounting, the verification suite that pins the mechanics, and the outputs.")
rep("Fig. 2. Interactive workspace: benchmark library (left) and a completed RC-column run with constraint tables and the translucent-concrete reinforcement view (right).",
    "Fig. 2. The workspace in action: (a) benchmark library; (b) live run with editable design inputs; (c) results with the governing check and the carbon breakdown; (d) optimized reinforcement cage; (e) AISC section properties and design details; (f) frequency convergence toward the target.")
rep("Figure 1 shows the architecture.", "Figure 1 shows the system flow.")
rep("This paper presents an open, interactive framework that unifies both objectives",
    "This paper presents an open, interactive educational framework that unifies both objectives")
i = find("reproducible from the accompanying open implementation.")
P[i].runs[-1].text += " The framework is designed first as a capacity-building tool: every run\u2014steel, aluminium or concrete\u2014exposes its governing limit state, its full section properties and an embodied-carbon total with breakdown, so learners connect design decisions to code checks and to carbon."
rep("steel frames; frequency constraints; engineering education",
    "steel frames; frequency constraints; engineering education; capacity building")
i = find("Because the factor enters the objective directly")
P[i].runs[-1].text += " Carbon is reported for every experiment, not only the RC family: steel and aluminium problems keep mass as the objective and report the embodied-carbon total with a breakdown by member role (or by flange and web) at a selectable material route\u2014structural steel defaulting to the world-average 1.85 kg/kg and the classic aluminium trusses to a European-mix 6.7 kg CO\u2082e/kg (indicative values from the International Aluminium Institute and the ICE database [28, 29]: 16.7 for world-average primary, 0.6 for recycled)."
i = find("Table 3 first summarizes one reproducible sample run per problem")
P[i].runs[0].text = P[i].runs[0].text.replace("Table 3 first summarizes one reproducible sample run per problem across the entire suite",
    "Table 3 first summarizes one reproducible sample run per problem across the entire suite, each with its embodied carbon at the default material route")
i = find("Documented model divergences discussed in the text")
P[i].runs[-1].text += "  \u00b2 Embodied carbon at the default material route (structural steel 1.85, aluminium 6.7, rebar 2.82 kg CO\u2082/kg); for the RC rows the CO\u2082 value is the objective itself."
i = find("This paper presented a verified, open, interactive framework")
P[i].runs[0].text = P[i].runs[0].text.replace("This paper presented a verified, open, interactive framework",
    "This paper presented a verified, open, interactive educational framework")
# Table 3: add the CO2 column
import copy as _copy
co2 = {"10-bar truss": "15,805", "25-bar tower": "779", "72-bar tower": "1,676", "10-member frame": "5,951",
       "25-member frame": "6,996", "3-storey steel MRF": "14,928", "Built-up I-beam, ASD": "584",
       "Built-up I-beam, LRFD": "512", "10-bar (frequency)": "17,248", "72-bar (frequency)": "1,901",
       "Portal (frequency)": "7,765", "RC beam": "728", "RC column": "163", "RC footing": "1,158",
       "RC 3-storey frame": "5,302"}
for row in T[2].rows:
    key = row.cells[0].text.strip()
    tc = _copy.deepcopy(row.cells[-1]._tc); row._tr.append(tc)
    from docx.table import _Cell
    setcell(_Cell(tc, T[2]), "CO\u2082 (kg)\u00b2" if key == "Problem" else co2.get(key, ""))
i26 = find("[27] AISC (2016)")
from docx.text.paragraph import Paragraph as _Par
for txt in ["[28] International Aluminium Institute (2021). Greenhouse Gas Emissions Intensity: Primary Aluminium. London.",
            "[29] Hammond, G., Jones, C. (2019). Inventory of Carbon and Energy (ICE) Database v3.0. University of Bath."]:
    newp = _copy.deepcopy(P[i26]._p); P[i26]._p.addnext(newp)
    _Par(newp, P[i26]._parent).runs[0].text = txt
    for r in _Par(newp, P[i26]._parent).runs[1:]: r.text = ""
    i26 += 0
# fix order: the two references were inserted in reverse; swap texts if needed
PP = doc.paragraphs
a = next(i for i, p in enumerate(PP) if p.text.startswith("[28]"))
b = next(i for i, p in enumerate(PP) if p.text.startswith("[29]"))
if a > b:
    ta, tb = PP[a].runs[0].text, PP[b].runs[0].text
    PP[a].runs[0].text, PP[b].runs[0].text = tb, ta
# page economy: trim the tallest remaining figures slightly
for sh in sorted(doc.inline_shapes, key=lambda z: -z.height)[:6]:
    sh.width = Emu(int(sh.width*0.86)); sh.height = Emu(int(sh.height*0.86))
rep("is available as an open, dependency-light implementation (NumPy/SciPy/Flask) with a 26-test verification suite that pins every mechanical anchor used in this paper, and a six-module capacity-building curriculum.",
    "is available as an open, dependency-light implementation (NumPy/SciPy/Flask) with its 43-test verification suite and capacity-building curriculum.")
rep(" (one page per problem: properties, cross-sections and a reproducible sample run)", "")
rep("whole-life carbon including modules A\u2013D; and a controlled classroom study of the educational deployment.",
    "whole-life carbon; and a controlled classroom study.")
rep("The framework\u2019s structure\u2014formulation, engine, accounting, optimizer, interface, each independently testable\u2014was designed so that each of these is an extension, not a rewrite.",
    "The framework\u2019s structure\u2014formulation, engine, accounting, optimizer, interface\u2014was designed so that each is an extension, not a rewrite.")
for p in doc.paragraphs:                       # 4% tighter leading, body only
    ls = p.paragraph_format.line_spacing
    if isinstance(ls, float) and ls > 1.05: p.paragraph_format.line_spacing = ls*0.96
doc.save(_os.path.join(OUT, "Structural_Optimization_Manuscript.docx"))
print("revision 2 applied")
# =========================== revision 3: verified reference list with DOIs
REFS = {
 1: "[1] Schmit, L.A., Farshi, B. (1974). Some approximation concepts for structural synthesis. AIAA Journal, 12(5), 692\u2013699. https://doi.org/10.2514/3.49321",
 2: "[2] Holland, J.H. (1975). Adaptation in Natural and Artificial Systems. University of Michigan Press, Ann Arbor, MI.",
 3: "[3] Goldberg, D.E. (1989). Genetic Algorithms in Search, Optimization and Machine Learning. Addison-Wesley, Reading, MA.",
 4: "[4] Rajeev, S., Krishnamoorthy, C.S. (1992). Discrete optimization of structures using genetic algorithms. Journal of Structural Engineering, 118(5), 1233\u20131250. https://doi.org/10.1061/(ASCE)0733-9445(1992)118:5(1233)",
 5: "[5] Sedaghati, R. (2005). Benchmark case studies in structural design optimization using the force method. International Journal of Solids and Structures, 42(21\u201322), 5848\u20135871. https://doi.org/10.1016/j.ijsolstr.2005.03.030",
 6: "[6] Grandhi, R.V. (1993). Structural optimization with frequency constraints \u2014 a review. AIAA Journal, 31(12), 2296\u20132303.",
 7: "[7] Pezeshk, S., Camp, C.V., Chen, D. (2000). Design of nonlinear framed structures using genetic optimization. Journal of Structural Engineering, 126(3), 382\u2013388. https://doi.org/10.1061/(ASCE)0733-9445(2000)126:3(382)",
 8: "[8] Camp, C.V., Bichon, B.J., Stovall, S.P. (2005). Design of steel frames using ant colony optimization. Journal of Structural Engineering, 131(3), 369\u2013379. https://doi.org/10.1061/(ASCE)0733-9445(2005)131:3(369)",
 9: "[9] De\u011fertekin, S.O. (2008). Optimum design of steel frames using harmony search algorithm. Structural and Multidisciplinary Optimization, 36(4), 393\u2013401. https://doi.org/10.1007/s00158-007-0177-4",
 10: "[10] To\u011fan, V. (2012). Design of planar steel frames using Teaching\u2013Learning Based Optimization. Engineering Structures, 34, 225\u2013232. https://doi.org/10.1016/j.engstruct.2011.08.035",
 11: "[11] Dumonteil, P. (1992). Simple equations for effective length factors. Engineering Journal, AISC, 29(3), 111\u2013115.",
 12: "[12] AISC (1994). Load and Resistance Factor Design Specification for Structural Steel Buildings, 2nd ed. American Institute of Steel Construction, Chicago, IL.",
 13: "[13] ACI Committee 318 (2019). Building Code Requirements for Structural Concrete (ACI 318-19) and Commentary. American Concrete Institute, Farmington Hills, MI.",
 14: "[14] Camp, C.V., Pezeshk, S., Hansson, H. (2003). Flexural design of reinforced concrete frames using a genetic algorithm. Journal of Structural Engineering, 129(1), 105\u2013115. https://doi.org/10.1061/(ASCE)0733-9445(2003)129:1(105)",
 15: "[15] Lee, C., Ahn, J. (2003). Flexural design of reinforced concrete frames by genetic algorithm. Journal of Structural Engineering, 129(6), 762\u2013774. https://doi.org/10.1061/(ASCE)0733-9445(2003)129:6(762)",
 16: "[16] Pay\u00e1-Zaforteza, I., Yepes, V., Hospitaler, A., Gonz\u00e1lez-Vidosa, F. (2009). CO2-optimization of reinforced concrete frames by simulated annealing. Engineering Structures, 31(7), 1501\u20131508. https://doi.org/10.1016/j.engstruct.2009.02.034",
 17: "[17] Yepes, V., Gonz\u00e1lez-Vidosa, F., Alcal\u00e1, J., Villalba, P. (2012). CO2-optimization design of reinforced concrete retaining walls based on a VNS-threshold acceptance strategy. Journal of Computing in Civil Engineering, 26(3), 378\u2013386. https://doi.org/10.1061/(ASCE)CP.1943-5487.0000140",
 18: "[18] Camp, C.V., Huq, F. (2013). CO2 and cost optimization of reinforced concrete frames using a big bang\u2013big crunch algorithm. Engineering Structures, 48, 363\u2013372. https://doi.org/10.1016/j.engstruct.2012.09.004",
 19: "[19] Camp, C.V., Assadollahi, A. (2013). CO2 and cost optimization of reinforced concrete footings using a hybrid big bang\u2013big crunch algorithm. Structural and Multidisciplinary Optimization, 48(2), 411\u2013426. https://doi.org/10.1007/s00158-013-0897-6",
 20: "[20] de Medeiros, G.F., Kripka, M. (2014). Optimization of reinforced concrete columns according to different environmental impact assessment parameters. Engineering Structures, 59, 185\u2013194. https://doi.org/10.1016/j.engstruct.2013.10.045",
 21: "[21] Yeo, D., Potra, F.A. (2015). Sustainable design of reinforced concrete structures through CO2 emission optimization. Journal of Structural Engineering, 141(3), B4014002. https://doi.org/10.1061/(ASCE)ST.1943-541X.0000888",
 22: "[22] Institut de Tecnologia de la Construcci\u00f3 de Catalunya (ITeC). BEDEC database of construction elements and unit environmental data. Barcelona, itec.cat (accessed 2026).",
 23: "[23] World Steel Association (2024). Sustainability Indicators 2024 Report. Brussels, worldsteel.org.",
 24: "[24] IEEFA (2022). Steel Fact Sheet: carbon intensity of BF\u2013BOF, DRI\u2013EAF and scrap\u2013EAF routes. Institute for Energy Economics and Financial Analysis, ieefa.org.",
 25: "[25] Columbia Business School, Climate Knowledge Initiative (2023). Steel Sector Overview: production routes and emission intensities. New York.",
 26: "[26] New Steel Construction / Eurofer (2024). The carbon footprint of steel: primary and secondary route intensities by system expansion. British Constructional Steelwork Association, newsteelconstruction.com.",
 27: "[27] AISC (2016). Specification for Structural Steel Buildings (ANSI/AISC 360-16). American Institute of Steel Construction, Chicago, IL.",
 28: "[28] International Aluminium Institute (2021). Greenhouse Gas Emissions Intensity \u2014 Primary Aluminium (statistical report). London, international-aluminium.org.",
 29: "[29] Hammond, G., Jones, C. (2019). Inventory of Carbon and Energy (ICE) Database, version 3.0. University of Bath / Circular Ecology, circularecology.com.",
}
from docx.shared import Pt
n_set = 0
for p in doc.paragraphs:
    t = p.text.strip()
    if t.startswith("[") and "]" in t[:5]:
        k = int(t[1:t.index("]")])
        if k in REFS:
            for r in p.runs[1:]: r.text = ""
            p.runs[0].text = REFS[k]; p.runs[0].font.size = Pt(8.5)
            p.paragraph_format.space_after = Pt(1.2)
            p.paragraph_format.line_spacing = 1.02
            n_set += 1
print("references rewritten:", n_set)
doc.save(_os.path.join(OUT, "Structural_Optimization_Manuscript.docx"))

